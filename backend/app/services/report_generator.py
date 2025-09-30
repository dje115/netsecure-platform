"""
Report Generation Service
Generates professional security assessment reports
"""

import logging
from typing import Dict, List, Optional
from datetime import datetime
from io import BytesIO
from sqlalchemy.orm import Session
from app.models.scan import ScanSession
from app.models.device import Device
from app.models.vulnerability import Vulnerability

logger = logging.getLogger(__name__)


class ReportGenerator:
    """Generate professional security assessment reports"""
    
    def __init__(self):
        self.report_templates = {
            "executive": "Executive Summary",
            "technical": "Technical Report",
            "comprehensive": "Comprehensive Assessment",
            "compliance": "Compliance Report"
        }
    
    async def generate_report(
        self,
        scan_session: ScanSession,
        db: Session,
        report_type: str = "comprehensive",
        format: str = "html"
    ) -> Dict:
        """
        Generate a security assessment report
        
        Args:
            scan_session: Scan session to report on
            db: Database session
            report_type: Type of report (executive, technical, comprehensive)
            format: Output format (html, pdf, docx)
        
        Returns:
            Report data and metadata
        """
        
        logger.info(f"Generating {report_type} report for scan {scan_session.id}")
        
        # Gather data
        report_data = await self._gather_report_data(scan_session, db)
        
        # Generate based on format
        if format == "html":
            content = await self._generate_html_report(report_data, report_type)
        elif format == "pdf":
            content = await self._generate_pdf_report(report_data, report_type)
        elif format == "docx":
            content = await self._generate_docx_report(report_data, report_type)
        else:
            raise ValueError(f"Unsupported format: {format}")
        
        return {
            "scan_id": scan_session.id,
            "report_type": report_type,
            "format": format,
            "generated_at": datetime.utcnow().isoformat(),
            "content": content,
            "summary": report_data["summary"]
        }
    
    async def _gather_report_data(
        self,
        scan_session: ScanSession,
        db: Session
    ) -> Dict:
        """Gather all data needed for the report"""
        
        # Get devices
        devices = db.query(Device).filter(
            Device.scan_session_id == scan_session.id
        ).all()
        
        # Get vulnerabilities
        vulnerabilities = []
        for device in devices:
            device_vulns = db.query(Vulnerability).filter(
                Vulnerability.device_id == device.id
            ).all()
            vulnerabilities.extend(device_vulns)
        
        # Calculate statistics
        total_devices = len(devices)
        hvt_devices = [d for d in devices if d.is_hvt]
        
        vuln_by_severity = {
            "critical": len([v for v in vulnerabilities if v.severity == "critical"]),
            "high": len([v for v in vulnerabilities if v.severity == "high"]),
            "medium": len([v for v in vulnerabilities if v.severity == "medium"]),
            "low": len([v for v in vulnerabilities if v.severity == "low"])
        }
        
        # Risk distribution
        risk_distribution = {
            "critical": len([d for d in devices if d.risk_level == "critical"]),
            "high": len([d for d in devices if d.risk_level == "high"]),
            "medium": len([d for d in devices if d.risk_level == "medium"]),
            "low": len([d for d in devices if d.risk_level == "low"])
        }
        
        return {
            "scan": {
                "id": scan_session.id,
                "name": scan_session.name,
                "target_range": scan_session.target_range,
                "scan_type": scan_session.scan_type,
                "started_at": scan_session.started_at.isoformat() if scan_session.started_at else None,
                "completed_at": scan_session.completed_at.isoformat() if scan_session.completed_at else None,
                "duration": self._calculate_duration(scan_session)
            },
            "summary": {
                "total_devices": total_devices,
                "hvt_count": len(hvt_devices),
                "total_vulnerabilities": len(vulnerabilities),
                "vuln_by_severity": vuln_by_severity,
                "risk_distribution": risk_distribution,
                "overall_risk_score": scan_session.risk_score or 0
            },
            "devices": [self._device_to_dict(d) for d in devices],
            "hvt_devices": [self._device_to_dict(d) for d in hvt_devices],
            "vulnerabilities": [self._vuln_to_dict(v) for v in vulnerabilities],
            "critical_findings": [self._vuln_to_dict(v) for v in vulnerabilities if v.severity == "critical"]
        }
    
    def _calculate_duration(self, scan: ScanSession) -> str:
        """Calculate scan duration"""
        if not scan.started_at or not scan.completed_at:
            return "N/A"
        
        duration = scan.completed_at - scan.started_at
        hours = duration.seconds // 3600
        minutes = (duration.seconds % 3600) // 60
        
        if hours > 0:
            return f"{hours}h {minutes}m"
        return f"{minutes}m"
    
    def _device_to_dict(self, device: Device) -> Dict:
        """Convert device to dict"""
        return {
            "id": device.id,
            "ip_address": device.ip_address,
            "hostname": device.hostname,
            "device_type": device.device_type,
            "os_name": device.os_name,
            "is_hvt": device.is_hvt,
            "hvt_type": device.hvt_type,
            "risk_score": device.risk_score,
            "risk_level": device.risk_level,
            "ai_classification": device.ai_classification
        }
    
    def _vuln_to_dict(self, vuln: Vulnerability) -> Dict:
        """Convert vulnerability to dict"""
        return {
            "id": vuln.id,
            "cve_id": vuln.cve_id,
            "title": vuln.title,
            "severity": vuln.severity,
            "cvss_score": vuln.cvss_score,
            "exploitable": vuln.exploitable,
            "remediation": vuln.remediation
        }
    
    async def _generate_html_report(
        self,
        data: Dict,
        report_type: str
    ) -> str:
        """Generate HTML report"""
        
        html = f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>Security Assessment Report - {data['scan']['name']}</title>
    <style>
        body {{
            font-family: Arial, sans-serif;
            line-height: 1.6;
            max-width: 1200px;
            margin: 0 auto;
            padding: 20px;
            color: #333;
        }}
        h1 {{ color: #2c3e50; border-bottom: 3px solid #3498db; padding-bottom: 10px; }}
        h2 {{ color: #34495e; margin-top: 30px; }}
        .executive-summary {{
            background: #ecf0f1;
            padding: 20px;
            border-radius: 5px;
            margin: 20px 0;
        }}
        .metric {{
            display: inline-block;
            margin: 10px 20px 10px 0;
            padding: 15px;
            background: #fff;
            border-radius: 5px;
            box-shadow: 0 2px 5px rgba(0,0,0,0.1);
        }}
        .metric-value {{
            font-size: 36px;
            font-weight: bold;
            color: #3498db;
        }}
        .metric-label {{
            font-size: 14px;
            color: #7f8c8d;
        }}
        .critical {{ color: #e74c3c; }}
        .high {{ color: #e67e22; }}
        .medium {{ color: #f39c12; }}
        .low {{ color: #3498db; }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
        }}
        th, td {{
            padding: 12px;
            text-align: left;
            border-bottom: 1px solid #ddd;
        }}
        th {{
            background: #34495e;
            color: white;
        }}
        .hvt-badge {{
            background: #e74c3c;
            color: white;
            padding: 3px 8px;
            border-radius: 3px;
            font-size: 12px;
        }}
        .footer {{
            margin-top: 50px;
            padding-top: 20px;
            border-top: 1px solid #ddd;
            text-align: center;
            color: #7f8c8d;
            font-size: 12px;
        }}
    </style>
</head>
<body>
    <h1>Security Assessment Report</h1>
    
    <div class="executive-summary">
        <h2>Executive Summary</h2>
        <p><strong>Scan Name:</strong> {data['scan']['name']}</p>
        <p><strong>Target Range:</strong> {data['scan']['target_range']}</p>
        <p><strong>Scan Type:</strong> {data['scan']['scan_type']}</p>
        <p><strong>Completed:</strong> {data['scan']['completed_at']}</p>
        <p><strong>Duration:</strong> {data['scan']['duration']}</p>
    </div>
    
    <h2>Key Metrics</h2>
    <div class="metric">
        <div class="metric-value">{data['summary']['total_devices']}</div>
        <div class="metric-label">Total Devices</div>
    </div>
    <div class="metric">
        <div class="metric-value critical">{data['summary']['hvt_count']}</div>
        <div class="metric-label">High-Value Targets</div>
    </div>
    <div class="metric">
        <div class="metric-value">{data['summary']['total_vulnerabilities']}</div>
        <div class="metric-label">Vulnerabilities</div>
    </div>
    <div class="metric">
        <div class="metric-value">{data['summary']['overall_risk_score']}/100</div>
        <div class="metric-label">Risk Score</div>
    </div>
    
    <h2>Vulnerability Distribution</h2>
    <div class="metric">
        <div class="metric-value critical">{data['summary']['vuln_by_severity']['critical']}</div>
        <div class="metric-label">Critical</div>
    </div>
    <div class="metric">
        <div class="metric-value high">{data['summary']['vuln_by_severity']['high']}</div>
        <div class="metric-label">High</div>
    </div>
    <div class="metric">
        <div class="metric-value medium">{data['summary']['vuln_by_severity']['medium']}</div>
        <div class="metric-label">Medium</div>
    </div>
    <div class="metric">
        <div class="metric-value low">{data['summary']['vuln_by_severity']['low']}</div>
        <div class="metric-label">Low</div>
    </div>
    
    <h2>High-Value Targets</h2>
    <table>
        <thead>
            <tr>
                <th>IP Address</th>
                <th>Hostname</th>
                <th>Type</th>
                <th>OS</th>
                <th>Risk Level</th>
            </tr>
        </thead>
        <tbody>
"""
        
        for device in data['hvt_devices']:
            html += f"""
            <tr>
                <td>{device['ip_address']}</td>
                <td>{device['hostname'] or 'N/A'}</td>
                <td><span class="hvt-badge">{device['hvt_type'] or 'HVT'}</span></td>
                <td>{device['os_name'] or 'Unknown'}</td>
                <td class="{device['risk_level']}">{device['risk_level'] or 'N/A'}</td>
            </tr>
"""
        
        html += """
        </tbody>
    </table>
    
    <h2>Critical Vulnerabilities</h2>
    <table>
        <thead>
            <tr>
                <th>CVE ID</th>
                <th>Title</th>
                <th>CVSS Score</th>
                <th>Exploitable</th>
            </tr>
        </thead>
        <tbody>
"""
        
        for vuln in data['critical_findings']:
            html += f"""
            <tr>
                <td>{vuln['cve_id'] or 'N/A'}</td>
                <td>{vuln['title']}</td>
                <td class="critical">{vuln['cvss_score'] or 'N/A'}</td>
                <td>{'Yes' if vuln['exploitable'] else 'No'}</td>
            </tr>
"""
        
        html += f"""
        </tbody>
    </table>
    
    <div class="footer">
        <p>Security Assessment Platform | Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC</p>
        <p>This report is confidential and should only be shared with authorized personnel.</p>
    </div>
</body>
</html>
"""
        
        return html
    
    async def _generate_pdf_report(
        self,
        data: Dict,
        report_type: str
    ) -> bytes:
        """Generate PDF report using ReportLab"""
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.lib import colors
            
            buffer = BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=letter)
            story = []
            styles = getSampleStyleSheet()
            
            # Title
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#2c3e50')
            )
            story.append(Paragraph("Security Assessment Report", title_style))
            story.append(Spacer(1, 0.3*inch))
            
            # Executive Summary
            story.append(Paragraph("Executive Summary", styles['Heading2']))
            summary_data = [
                ["Scan Name:", data['scan']['name']],
                ["Target Range:", data['scan']['target_range']],
                ["Completed:", data['scan']['completed_at']],
                ["Total Devices:", str(data['summary']['total_devices'])],
                ["HVT Count:", str(data['summary']['hvt_count'])],
                ["Vulnerabilities:", str(data['summary']['total_vulnerabilities'])],
                ["Risk Score:", f"{data['summary']['overall_risk_score']}/100"]
            ]
            summary_table = Table(summary_data)
            summary_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ]))
            story.append(summary_table)
            
            # Build PDF
            doc.build(story)
            return buffer.getvalue()
            
        except ImportError:
            logger.warning("ReportLab not installed, returning HTML instead")
            html = await self._generate_html_report(data, report_type)
            return html.encode('utf-8')
    
    async def _generate_docx_report(
        self,
        data: Dict,
        report_type: str
    ) -> bytes:
        """Generate DOCX report using python-docx"""
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            
            doc = Document()
            
            # Title
            title = doc.add_heading('Security Assessment Report', 0)
            title.alignment = 1  # Center
            
            # Executive Summary
            doc.add_heading('Executive Summary', 1)
            doc.add_paragraph(f"Scan Name: {data['scan']['name']}")
            doc.add_paragraph(f"Target Range: {data['scan']['target_range']}")
            doc.add_paragraph(f"Completed: {data['scan']['completed_at']}")
            doc.add_paragraph(f"Duration: {data['scan']['duration']}")
            
            # Metrics
            doc.add_heading('Key Metrics', 1)
            metrics_table = doc.add_table(rows=5, cols=2)
            metrics_table.style = 'Light Grid Accent 1'
            
            metrics_data = [
                ['Total Devices', str(data['summary']['total_devices'])],
                ['High-Value Targets', str(data['summary']['hvt_count'])],
                ['Total Vulnerabilities', str(data['summary']['total_vulnerabilities'])],
                ['Risk Score', f"{data['summary']['overall_risk_score']}/100"],
                ['Critical Vulnerabilities', str(data['summary']['vuln_by_severity']['critical'])]
            ]
            
            for i, (label, value) in enumerate(metrics_data):
                metrics_table.rows[i].cells[0].text = label
                metrics_table.rows[i].cells[1].text = value
            
            # Save to buffer
            buffer = BytesIO()
            doc.save(buffer)
            return buffer.getvalue()
            
        except ImportError:
            logger.warning("python-docx not installed, returning HTML instead")
            html = await self._generate_html_report(data, report_type)
            return html.encode('utf-8')


# Global report generator instance
report_generator = ReportGenerator()
